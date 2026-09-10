from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import chromadb
from sentence_transformers import SentenceTransformer
from langchain_groq import ChatGroq
from dotenv import load_dotenv
# pyrefly: ignore [missing-import]
from pypdf import PdfReader
import docx
import uvicorn
import io
import os

load_dotenv()

app = FastAPI(title="Academic AI - RAG Engine")

# CORS — allows frontend to call this API from any origin
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Lock this down in production
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# In-memory vector DB
chroma_client = chromadb.Client() 
collection = chroma_client.get_or_create_collection(name="fast_notes")

embedder = SentenceTransformer("BAAI/bge-small-en-v1.5")
llm = ChatGroq(
    model="qwen/qwen3.8-27b",
    temperature=0.1,
    api_key=os.getenv("GROQ_API_KEY")
)

class QueryRequest(BaseModel):
    query: str

@app.post("/ingest")
async def ingest_file(file: UploadFile = File(...)):
    filename = file.filename.lower()
    content = await file.read()
    
    if filename.endswith(".txt"):
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            text = content.decode("latin-1", errors="ignore")
    elif filename.endswith(".pdf"):
        try:
            pdf_reader = PdfReader(io.BytesIO(content))
            extracted_pages = [page.extract_text() or "" for page in pdf_reader.pages]
            text = "\n".join(extracted_pages).strip()
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to read PDF file: {str(e)}")
    elif filename.endswith(".docx"):
        try:
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n".join(paragraphs).strip()
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to read DOCX file: {str(e)}")
    else:
        raise HTTPException(status_code=400, detail="Only .txt, .pdf, and .docx files allowed.")
    
    if not text.strip():
        raise HTTPException(status_code=400, detail="Uploaded file contains no readable text.")
    
    chunk_size = 500
    chunks = [text[i:i+chunk_size] for i in range(0, len(text), chunk_size)]
    
    embeddings = embedder.encode(chunks).tolist()
    ids = [f"{file.filename}_{i}" for i in range(len(chunks))]
    metadatas = [{"source": file.filename} for _ in chunks]
    
    collection.add(ids=ids, documents=chunks, embeddings=embeddings, metadatas=metadatas)
    return {"status": "success", "filename": file.filename, "chunks_loaded": len(chunks)}

import re

def clean_text(text: str) -> str:
    """Removes markdown bold/italic asterisks, backticks, hashtags, and cleans whitespace."""
    text = re.sub(r'\*{1,3}(.*?)\*{1,3}', r'\1', text)
    text = re.sub(r'_{1,3}(.*?)_{1,3}', r'\1', text)
    text = re.sub(r'^#{1,6}\s*', '', text, flags=re.MULTILINE)
    text = re.sub(r'`(.*?)`', r'\1', text)
    text = re.sub(r'^\s*[\*\-]\s+', '• ', text, flags=re.MULTILINE)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

@app.post("/ask")
async def ask_question(payload: QueryRequest):
    query_vector = embedder.encode([payload.query]).tolist()
    results = collection.query(query_embeddings=query_vector, n_results=2)
    
    if not results["documents"][0]:
        return {"status": "escalated_to_teacher", "reason": "No notes uploaded yet."}
        
    distance = results["distances"][0][0]
    
    if distance > 0.7:
         return {"status": "escalated_to_teacher", "reason": "Low confidence match."}
         
    context = "\n".join(results["documents"][0])
    prompt = (
        "You are an academic teaching assistant. Answer the question strictly using only the context provided below.\n"
        "Guidelines:\n"
        "1. If the provided context is incomplete or does not contain enough information to fully answer the question, do NOT speculate or produce an incomplete response. Instead, start your response with 'ESCALATE:' followed by a clear, concise explanation of what information is missing.\n"
        "2. Write in clean, plain academic English. Avoid using markdown formatting tags, asterisks (**), hashtags (#), or code backticks (`). Present your answer in clear, well-structured paragraphs.\n\n"
        f"Context:\n{context}\n\n"
        f"Question: {payload.query}\n\n"
        "Answer:"
    )
    
    response = llm.invoke(prompt)
    raw_answer = response.content.strip()
    
    # Check if the model flagged insufficient context or requested escalation
    lower_answer = raw_answer.lower()
    is_escalated = (
        raw_answer.upper().startswith("ESCALATE:")
        or "cannot be derived from the given" in lower_answer
        or "not possible to" in lower_answer
        or "insufficient information" in lower_answer
    )
    
    if is_escalated:
        reason = raw_answer
        if reason.upper().startswith("ESCALATE:"):
            reason = reason[len("ESCALATE:"):].strip()
        return {
            "status": "escalated_to_teacher",
            "reason": clean_text(reason),
            "citations": [results["metadatas"][0][0]["source"]]
        }
    
    return {
        "status": "answered",
        "answer": clean_text(raw_answer),
        "citations": [results["metadatas"][0][0]["source"]]
    }

if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=8000)

